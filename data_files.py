# data_files.py
# -------------------------------------------------------------------
# Builder genérico de archivos de datos para solicitudes internas
# ("X File (Excel needed)"). NO es específico de "Best Buy": arma un
# Excel limpio a partir de tablas que el agente estructura con los datos
# que el remitente provee (Opción A). El builder NO inventa datos.
#
# API:
#   generar_xlsx_tabla(spec, output_path) -> str
#
# spec admite DOS formas:
#  (1) Una sola hoja:
#      {"title","sheet_name","headers":[...],"rows":[[...],...]}
#  (2) Varias hojas (workbook de varias pestañas) — para pedidos tipo
#      "8-tab workbook":
#      {"filename": "...",
#       "sheets": [
#          {"sheet_name":"Cover","title":"...","headers":[...],"rows":[...]},
#          {"sheet_name":"Route Summary","headers":[...],"rows":[...]},
#          ... ]}
# -------------------------------------------------------------------

NAVY_HEX = "1F3A5F"
LIGHT_HEX = "EEF2F7"
DARK_HEX = "222222"


def _nombre_hoja(propuesto, usados, idx):
    base = (str(propuesto or "").strip() or f"Sheet{idx}")[:31]
    nombre = base
    n = 2
    while nombre.lower() in usados:
        suf = f"_{n}"
        nombre = (base[:31 - len(suf)] + suf)
        n += 1
    usados.add(nombre.lower())
    return nombre


def _escribir_hoja(ws, hoja):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    hoja = hoja or {}
    headers = [str(h) for h in (hoja.get("headers") or [])]
    rows = hoja.get("rows") or []
    title = str(hoja.get("title") or "").strip()

    if not headers and rows:
        ncols = max((len(r) for r in rows if isinstance(r, (list, tuple))), default=1)
        headers = [f"Column {i+1}" for i in range(ncols)]
    ncols = max(len(headers), 1)

    navy_fill = PatternFill("solid", fgColor=NAVY_HEX)
    light_fill = PatternFill("solid", fgColor=LIGHT_HEX)
    white_bold = Font(bold=True, color="FFFFFF", size=10)
    dark = Font(color=DARK_HEX, size=10)
    thin = Side(style="thin", color="CCD4DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    r = 1
    if title:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        c = ws.cell(row=r, column=1, value=title)
        c.font = Font(bold=True, color=NAVY_HEX, size=13)
        r += 2

    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=j, value=h)
        cell.fill = navy_fill
        cell.font = white_bold
        cell.border = border
        cell.alignment = left
    header_row = r
    r += 1

    for idx, fila in enumerate(rows, start=1):
        if not isinstance(fila, (list, tuple)):
            fila = [fila]
        valores = list(fila)[:ncols] + [""] * (ncols - len(fila))
        for j, v in enumerate(valores, start=1):
            cell = ws.cell(row=r, column=j, value=("" if v is None else str(v)))
            cell.font = dark
            cell.border = border
            cell.alignment = left
            if idx % 2 == 0:
                cell.fill = light_fill
        r += 1

    for j in range(1, ncols + 1):
        largo = len(str(headers[j - 1])) if j - 1 < len(headers) else 8
        for fila in rows:
            if isinstance(fila, (list, tuple)) and j - 1 < len(fila):
                largo = max(largo, len(str(fila[j - 1])))
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(largo + 3, 48))

    ws.sheet_view.showGridLines = False
    if rows:
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(ncols)}{header_row + len(rows)}"
        )
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def generar_xlsx_tabla(spec, output_path):
    from openpyxl import Workbook

    spec = spec or {}
    wb = Workbook()
    hojas = spec.get("sheets")

    if hojas and isinstance(hojas, list):
        usados = set()
        primera = True
        for idx, h in enumerate(hojas, start=1):
            ws = wb.active if primera else wb.create_sheet()
            ws.title = _nombre_hoja((h or {}).get("sheet_name") or (h or {}).get("title"),
                                    usados, idx)
            _escribir_hoja(ws, h)
            primera = False
        if primera:  # sheets venía vacío -> hoja única desde el spec
            wb.active.title = _nombre_hoja(spec.get("sheet_name"), set(), 1)
            _escribir_hoja(wb.active, spec)
    else:
        wb.active.title = _nombre_hoja(spec.get("sheet_name"), set(), 1)
        _escribir_hoja(wb.active, spec)

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    demo = {
        "filename": "Target_Travel_Budget.xlsx",
        "sheets": [
            {"sheet_name": "Cover", "title": "Target Routes — Travel Budget",
             "headers": ["Field", "Value"],
             "rows": [["Prepared for", "Richard"], ["Scope", "Pricing only, no bookings"]]},
            {"sheet_name": "Route Summary",
             "headers": ["Crew", "State", "Size", "Stores", "Start", "End"],
             "rows": [["40", "CA", "5", "6", "2026-07-13", "2026-07-19"],
                      ["29", "TX", "6", "6", "2026-07-13", "2026-07-19"]]},
            {"sheet_name": "Texas Airbnb",
             "headers": ["Crew", "City", "Nightly", "Nights", "Total", "Link"],
             "rows": [["29", "Longview TX", "$140", "6", "$840", "verify"]]},
        ],
    }
    print("XLSX:", generar_xlsx_tabla(demo, "Target_multi_sample.xlsx"))
